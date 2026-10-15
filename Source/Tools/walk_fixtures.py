"""The files a walk of the built window works on (Source\\Tools\\release-walk.ps1), made by the Python of the stand.

    python walk_fixtures.py <folder> scan [count]   pages of a scan: text and a ruled table of numbers, PNG
    python walk_fixtures.py <folder> docx           an office document: a heading, a paragraph, a table
    python walk_fixtures.py <folder> md             Markdown: headings, a paragraph, a table

Nothing here is read from the owner's files; the text is made up for the walk.
"""

from __future__ import annotations

import sys
from pathlib import Path

LINES = [
    "ПОЯСНИТЕЛЬНАЯ ЗАПИСКА",
    "Проект планировки территории квартала 12",
    "Explanatory note: coordinates of the boundary points",
]
TABLE = [
    ["N", "X", "Y"],
    ["1", "452 118,25", "2 217 904,10"],
    ["2", "452 163,40", "2 217 951,75"],
    ["3", "452 129,05", "2 218 002,30"],
]


def font(size: int):
    from PIL import ImageFont

    for name in (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\segoeui.ttf"):
        if Path(name).is_file():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def scan(folder: Path, count: int) -> None:
    from PIL import Image, ImageDraw

    big, small = font(44), font(36)
    for index in range(count):
        page = Image.new("L", (1654, 2339), 255)
        draw = ImageDraw.Draw(page)
        y = 160
        for line in LINES:
            draw.text((140, y), line, fill=0, font=big if y == 160 else small)
            y += 90
        top, left, widths, height = y + 60, 140, [180, 520, 520], 90
        for row, cells in enumerate(TABLE):
            x = left
            for column, cell in enumerate(cells):
                draw.rectangle([x, top + row * height, x + widths[column], top + (row + 1) * height], outline=0, width=3)
                draw.text((x + 20, top + row * height + 22), cell, fill=0, font=small)
                x += widths[column]
        # The lines Workbench takes for coordinates: two numbers and a mark of an axis on one line.
        y = top + len(TABLE) * height + 80
        for row in TABLE[1:]:
            draw.text((140, y), f"Точка {row[0]}: X {row[1].replace(' ', '').replace(',', '.')} Y {row[2].replace(' ', '').replace(',', '.')}", fill=0, font=small)
            y += 70
        draw.text((140, y + 40), f"Page {index + 1}", fill=0, font=small)
        page.save(folder / f"scan_{index + 1:03d}.png", dpi=(200, 200))


def docx(folder: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading(LINES[0], level=1)
    document.add_paragraph(LINES[1])
    table = document.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    for row, cells in enumerate(TABLE):
        for column, cell in enumerate(cells):
            table.cell(row, column).text = cell
    document.save(folder / "note.docx")


def markdown(folder: Path) -> None:
    rows = ["| " + " | ".join(cells) + " |" for cells in TABLE]
    rows.insert(1, "| " + " | ".join("---" for _ in TABLE[0]) + " |")
    text = "\n".join([f"# {LINES[0]}", "", LINES[1], "", "## Points", "", *rows, "", LINES[2], ""])
    (folder / "note.md").write_text(text, encoding="utf-8", newline="\n")


def main(argv: list[str]) -> int:
    folder, kind = Path(argv[1]), argv[2]
    folder.mkdir(parents=True, exist_ok=True)
    if kind == "scan":
        scan(folder, int(argv[3]) if len(argv) > 3 else 1)
    elif kind == "docx":
        docx(folder)
    elif kind == "md":
        markdown(folder)
    else:
        print(f"unknown kind: {kind}")
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
